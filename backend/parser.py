import pdfplumber
from docx import Document
import io
import re

# These terms must not be split by the camelCase fixer
PRESERVE_TERMS = [
    "FastAPI", "PyTorch", "TensorFlow", "MongoDB", "MySQL", "JavaScript",
    "TypeScript", "OpenCV", "NumPy", "GitHub", "GitLab", "LinkedIn",
    "YouTube", "HuggingFace", "Scikit", "PowerBI", "NodeJS", "ReactJS",
    "ExpressJS", "TailwindCSS", "VGG16", "ResNet", "Socket.io", "ChatGPT",
    "DeepLearning", "MachineLearning", "ComputerVision"
]

def fix_spacing(text: str) -> str:
    # Temporarily replace known terms with placeholders
    placeholders = {}
    for i, term in enumerate(PRESERVE_TERMS):
        placeholder = f"__TERM{i}__"
        placeholders[placeholder] = term
        text = text.replace(term, placeholder)

    # Insert space before capital letters that follow lowercase
    text = re.sub(r'([a-z])([A-Z])', r'\1 \2', text)
    # Fix merged words after punctuation
    text = re.sub(r'([,;:])([A-Za-z])', r'\1 \2', text)
    # Collapse multiple spaces
    text = re.sub(r' +', ' ', text)

    # Restore preserved terms
    for placeholder, term in placeholders.items():
        text = text.replace(placeholder, term)

    return text.strip()

def extract_text_from_pdf(file_bytes: bytes) -> str:
    text = ""
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
    return fix_spacing(text)

def extract_text_from_docx(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    text = "\n".join([para.text for para in doc.paragraphs])
    return fix_spacing(text)

def extract_text(file_bytes: bytes, filename: str) -> str:
    if filename.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif filename.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    else:
        return file_bytes.decode("utf-8", errors="ignore").strip()